#!/usr/bin/env python3
"""التحديث:  python update.py
بيقرا: icons*/ + Originals.xlsx (غير معدّل) + Products.xlsx (أكونتات) + Accessories.xlsx (ملحقات) + Genres.xlsx (تصنيف الألعاب) + config.json
وبيكتب: data/catalog.js
الكتالوج مستقل تمامًا: كل بياناته جوه الفولدر ومش بيقرا أي داتابيز."""
import re, json, difflib, pathlib
from openpyxl import load_workbook

R = pathlib.Path(__file__).parent
cfg = json.loads((R / 'config.json').read_text('utf-8'))
SEC = cfg['sections']
icons_of = lambda d: sorted(p.name for p in (R / d).glob('*.webp')) if (R / d).exists() else []
stem = lambda f: re.sub(r'\s*\(ps2\)$', '', pathlib.PurePosixPath(f).name[:-5], flags=re.I)
norm = lambda s: re.sub(r'[^0-9a-z\u0600-\u06ff]', '', s.lower())
digits = lambda s: re.findall(r'\d+', s)
by = {}
for d in reversed([sc['folder'] for sc in SEC] + cfg.get('extra_icon_folders', [])):  # لو نفس اللعبة في أكتر من فولدر، أول قسم هو اللي يكسب
    by.update({norm(stem(f)): d + '/' + f for f in icons_of(d)})
alias = {k.lower(): v for k, v in cfg['aliases'].items()}
VAR = {'e': 'English', 'ar': 'عربي', 'عربي': 'عربي', 'ultimate': 'Ultimate', 'h': '', 'only': ''}
P45 = r'ps\s*\(?4\s*\+\s*(ps)?\s*5\)?'


def parse(raw):
    s, badges = raw, []
    if re.search(P45, s, re.I):
        s = re.sub(P45, ' ', s, flags=re.I); badges.append('PS4 + PS5')
    if re.search(r'ps5\s*only', s, re.I):
        s = re.sub(r'ps5\s*only', ' ', s, flags=re.I); badges.append('PS5 فقط')
    elif re.search(r'\bps5\b', s, re.I):
        s = re.sub(r'\bps5\b', ' ', s, flags=re.I); badges.append('PS5')
    out = []
    for p in (x.strip() for x in s.split('+') if x.strip()):
        w = p.split()
        while len(w) > 1 and w[-1].lower() in VAR:
            v = VAR[w.pop().lower()]
            if v and v not in badges: badges.append(v)
        out.append(' '.join(w))
    return out, badges


def match(c):
    key = norm(alias.get(c.lower(), c))
    if key in by: return by[key]
    for k in difflib.get_close_matches(key, list(by), 3, 0.86):
        if digits(k) == digits(key): return by[k]


# ---- تصنيف الألعاب: Genres.xlsx (A اللعبة | B التصنيف، وممكن أكتر من تصنيف بعلامة +)
def xrows(name):
    f = R / name
    return list(load_workbook(f, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True)) if f.exists() else []


def ytid(v):  # يقبل لينك يوتيوب بأي شكل أو الـID لوحده
    m = re.search(r'(?:v=|youtu\.be/|embed/|shorts/)([\w-]{11})', v) or re.fullmatch(r'([\w-]{11})', v.strip())
    return m.group(1) if m else ''


gmap, bmap, vmap, pmap, gorder, grows = {}, {}, {}, {}, [], []
for row in xrows('Genres.xlsx'):
    name = str(row[0] or '').strip()
    g = str(row[1] or '').strip() if len(row) > 1 else ''
    bl = str(row[2] or '').strip() if len(row) > 2 else ''  # C = نبذة
    if not name: continue
    f = match(name)
    key = norm(stem(f)) if f else norm(name)
    vd = str(row[3] or '').strip() if len(row) > 3 else ''  # D = فيديو (لينك يوتيوب أو ID)
    pl = str(row[4] or '').strip() if len(row) > 4 else ''  # E = عدد اللاعبين
    if pl: pmap[key] = pl
    if bl: bmap[key] = bl
    if vd and ytid(vd): vmap[key] = ytid(vd)
    if not g: continue
    gs = [x.strip() for x in re.split(r'[+،,]', g) if x.strip()]
    gmap[key] = gs
    grows.append((name, f))
    gorder += [x for x in gs if x not in gorder]
pref = cfg.get('genre_order', [])
gorder = [g for g in pref if g in gorder] + [g for g in gorder if g not in pref]
gen_of = lambda title: gmap.get(norm(title), [])
bl_of = lambda title: bmap.get(norm(title), '')
vd_of = lambda title: vmap.get(norm(title), '')
pl_of = lambda title: pmap.get(norm(title), '')

# ---- ألعاب المعدّل: كل قسم = فولدر أيقونات + سعر (config.json ← sections)
games = {}
for sc in SEC:
    games[sc['id']] = [{'title': stem(f), 'icon': sc['folder'] + '/' + f, 'ps2': bool(re.search(r'\(ps2\)$', f[:-5], re.I)),
                        'price': cfg['modded_prices'].get(stem(f), sc['price']), 'genres': gen_of(stem(f)), 'blurb': bl_of(stem(f)), 'video': vd_of(stem(f)), 'players': pl_of(stem(f))}
                       for f in icons_of(sc['folder'])]

# ---- الأكونتات: من الإكسيل (العمود B اختياري = السعر)
accounts, missing = [], set()
for row in load_workbook(R / 'Products.xlsx', read_only=True).active.iter_rows(min_row=2, values_only=True):
    raw = (row[0] or '').strip() if row else ''
    if not raw: continue
    comps, badges = parse(raw)
    pb = [b for b in badges if b in ('PS4 + PS5', 'PS5 فقط', 'PS5')]
    badges = [b for b in badges if b not in pb]
    cell = str(row[2] or '').upper() if len(row) > 2 else ''  # C = المنصة: PS4 / PS5 / PS4 + PS5
    if '4' in cell or '5' in cell: plats = [x for x, k in (('PS4', '4' in cell), ('PS5', '5' in cell)) if k]
    elif 'PS4 + PS5' in pb: plats = ['PS4', 'PS5']
    elif pb: plats = ['PS5']
    else: plats = ['PS4']
    if len(plats) == 2: badges.insert(0, 'PS4 + PS5')
    items = []
    for c in comps:
        f = match(c)
        if not f: missing.add(c)
        items.append({'title': stem(f) if f else ' '.join(w[:1].upper() + w[1:] for w in c.split()), 'icon': f, 'blurb': bl_of(stem(f)) if f else '', 'video': vd_of(stem(f)) if f else vd_of(c), 'players': pl_of(stem(f)) if f else pl_of(c)})
    price = row[1] if len(row) > 1 and isinstance(row[1], (int, float)) else cfg['account_price']
    accounts.append({'name': ' + '.join(i['title'] for i in items), 'items': items, 'badges': badges, 'price': price, 'platforms': plats})

AP = cfg['account_tabs']  # الأكونت المشترك بيظهر في التبويبين
accts = {t['id']: [a for a in accounts if t['platform'] in a['platforms']] for t in AP}

# ---- غير معدّل: Originals.xlsx (A اللعبة | B Off | C Online | D Pr | E المنصة اختياري)
oc = cfg['originals']
originals = []
for row in xrows(oc['file']):
    n = str(row[0] or '').strip()
    if not n: continue
    f = match(n)
    prices = [[t, v] for t, v in zip(oc['types'], row[1:4]) if isinstance(v, (int, float)) and v]
    plat = str(row[4] or '').strip() if len(row) > 4 else ''
    originals.append({'title': n, 'icon': f, 'prices': prices, 'badges': [plat] if plat else [],
                      'blurb': bl_of(n) or (bl_of(stem(f)) if f else ''), 'genres': gen_of(n) or (gen_of(stem(f)) if f else []), 'video': vd_of(n) or (vd_of(stem(f)) if f else ''), 'players': pl_of(n) or (pl_of(stem(f)) if f else '')})

# ---- الملحقات: Accessories.xlsx (A الاسم | B السعر | C متوفر نعم/لا | D التصنيف اختياري)
CATS = [('شاحن', 'شواحن'), ('وصلة', 'وصلات وكابلات'), ('كبل', 'وصلات وكابلات'), ('استند', 'استاندات'),
        ('جراب', 'جرابات وسكنات'), ('سكن', 'جرابات وسكنات'), ('شنط', 'جرابات وسكنات'), ('دراع', 'دراعات')]
NO = {'لا', 'no', 'n', '0', 'نفذ', 'نفذت', 'غير متوفر', 'false'}
products = []
for row in xrows('Accessories.xlsx'):
    n = str(row[0] or '').strip()
    if not n: continue
    price = row[1] if len(row) > 1 and isinstance(row[1], (int, float)) and row[1] else None
    st = row[2] if len(row) > 2 else None
    ok = st > 0 if isinstance(st, (int, float)) else str(st or '').strip().lower() not in NO
    cat = (str(row[3] or '').strip() if len(row) > 3 else '') or next((c for k, c in CATS if k in n), 'ملحقات أخرى')
    products.append({'name': n, 'price': price, 'cat': cat, 'in_stock': ok})

tabs = [{'id': sc['id'], 'title': sc['title'], 'price': sc['price'], 'symbol': sc.get('symbol', '△'), 'color': sc.get('color', '--tri'), 'kind': 'modded'}
        for sc in SEC if games[sc['id']]]
if originals: tabs.append({'id': 'originals', 'title': oc['title'], 'symbol': '□', 'color': '--sq', 'kind': 'originals'})
tabs += [{'id': t['id'], 'title': t['title'], 'symbol': '○', 'color': '--cir', 'kind': 'accounts'} for t in AP if accts[t['id']]]
tabs.append({'id': 'products', 'title': 'ملحقات', 'symbol': '✕', 'color': '--crs', 'kind': 'products'})
out = {'shop': cfg['shop'], 'phones': cfg['phones'], 'tabs': tabs, 'games': games, 'originals': originals,
       'accts': accts, 'products': products, 'genres': gorder, 'orig': {'title': oc['title'], 'types': oc['types']}}
(R / 'data').mkdir(exist_ok=True)
(R / 'data' / 'catalog.js').write_text('window.CATALOG=' + json.dumps(out, ensure_ascii=False) + ';', 'utf-8')

# ---- صفحات المشاركة: g/<slug>.html بتحمل صورة اللعبة (Open Graph) عشان الواتساب يعرض غلاف اللعبة بدل اللوجو
import struct, html as _h, urllib.parse as _u
from PIL import Image, ImageFilter
SITE = cfg.get('site_url', 'https://zaza-184.github.io/').rstrip('/') + '/'
(R / 'g').mkdir(exist_ok=True); (R / 'og').mkdir(exist_ok=True)
def slug(t, k):  # لازم يطابق slug في index.html و admin.html
    b = (t + '|' + k).encode('utf-16-le'); h = 2166136261
    for u in struct.unpack('<%dH' % (len(b) // 2), b): h = ((h ^ u) * 16777619) & 0xFFFFFFFF
    return (re.sub(r'[^a-z0-9]+', '-', k.lower()).strip('-')[:40] or 'g') + '-' + format(h, 'x')
def og(icon):
    if not icon: return 'assets/og.jpg'
    out = 'og/' + slug('i', icon) + '.jpg'
    if not (R / out).exists() or (R / out).stat().st_mtime < (R / icon).stat().st_mtime:
        im = Image.open(R / icon).convert('RGB')
        bg = im.resize((1200, 1200)).crop((0, 285, 1200, 915)).filter(ImageFilter.GaussianBlur(28)).point(lambda v: int(v * .45))
        bg.paste(im.resize((560, 560)), (320, 35)); bg.save(R / out, 'JPEG', quality=80)
    return out
try: LV = json.loads((R / 'data' / 'live.json').read_text('utf-8'))
except Exception: LV = {}
keep = {(g['tab'], g['title']) for g in LV.get('added', [])} | {(t, k) for t, m in LV.get('edit', {}).items() for k, e in m.items() if any(f in e for f in ('img', 'dn', 'items'))}
DESC = 'اطلب دلوقتي من ZAZA بتاع ألعاب على واتساب'
def page(tab, key, desc, icon):
    if (tab, key) in keep: return
    sl, kq, q = slug(tab, key), _u.quote(key, safe=''), _h.escape
    (R / 'g' / (sl + '.html')).write_text(f'<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8"><title>{q(key)} | ZAZA</title><meta property="og:type" content="website"><meta property="og:title" content="{q(key)} | ZAZA بتاع ألعاب"><meta property="og:description" content="{q(desc or DESC)}"><meta property="og:image" content="{SITE}{og(icon)}"><meta property="og:image:width" content="1200"><meta property="og:image:height" content="630"><meta property="og:url" content="{SITE}g/{sl}.html"><meta name="twitter:card" content="summary_large_image"><script>location.replace("../#{tab}/{kq}")</script><a href="../#{tab}/{kq}">افتح</a></html>', 'utf-8')
for tid, l in games.items():
    for g in l: page(tid, g['title'], g['blurb'], g['icon'])
for g in originals: page('originals', g['title'], g['blurb'], g['icon'])
for tid, l in accts.items():
    for a in l: page(tid, a['name'], next((i['blurb'] for i in a['items'] if len(a['items']) == 1 and i['blurb']), ''), next((i['icon'] for i in a['items'] if i['icon']), None))
print('صفحات المشاركة: g/ و og/ اتحدّثت')
print(' | '.join(f"{sc['title']}: {len(games[sc['id']])}" for sc in SEC) +
      f' | {oc["title"]}: {len(originals)} | ' + ' | '.join(f"{t['title']}: {len(accts[t['id']])}" for t in AP) + f' | ملحقات: {len(products)}')
print('ألعاب في الأكونتات من غير أيقونة (ضيف الأيقونة في icons/ أو alias في config.json):')
print('  ' + ' | '.join(sorted(missing)))
allg = [i for l in [*games.values(), originals] for i in l]
nog = [i['title'] for i in allg if not i['genres']]
if nog: print('ألعاب من غير تصنيف (ضيفها في Genres.xlsx):\n  ' + ' | '.join(nog))
titles = {norm(i['title']) for i in allg}
orphan = [n for n, f in grows if not f and norm(n) not in titles]
if orphan: print('أسماء في Genres.xlsx مالهاش لعبة (راجع الاسم):\n  ' + ' | '.join(orphan))
noic = [i['title'] for i in originals if not i['icon']]
if noic: print('ألعاب غير معدّل من غير أيقونة (ضيفها في أي فولدر icons):\n  ' + ' | '.join(noic))
