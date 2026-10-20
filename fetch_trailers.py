#!/usr/bin/env python3
"""بيملّا عمود "فيديو" (D) في Genres.xlsx بأفضل تريلر رسمي لكل لعبة، عن طريق YouTube Data API.

الاستخدام:
    python fetch_trailers.py YOUR_API_KEY
    (أو حط المفتاح في متغير بيئة اسمه YT_KEY)

- بيتخطى أي لعبة فيها فيديو بالفعل (فبتقدر تعدّل بإيدك وهو مش هيلمسها).
- بيختار بس فيديوهات مسموح تشغيلها جوه المواقع (videoEmbeddable) عشان المشغّل يشتغل.
- كل بحث بياخد 100 وحدة من الكوتا المجانية (10,000 في اليوم ≈ 100 لعبة)، فلما الكوتا تخلص بيقف
  ويحفظ اللي خلّصه، وشغّله تاني بكرة يكمّل من حيث وقف.
- بيكتب trailers_log.csv فيه اسم الفيديو والقناة لكل لعبة عشان تراجع الاختيارات.
- بعدها شغّل:  python update.py
"""
import csv, json, os, re, sys, urllib.error, urllib.parse, urllib.request
from pathlib import Path
from openpyxl import load_workbook

R = Path(__file__).parent
BAD = ('review', 'reaction', 'walkthrough', "let's play", 'lets play', 'tutorial', 'how to', 'part 1', 'ending',
       'all cutscenes', 'speedrun', ' mod', 'unboxing', 'tier list', ' vs ', 'comparison', 'top 10', 'fan made')


def score(title, game):
    t = title.lower()
    words = [w for w in re.sub(r'[^a-z0-9 ]', ' ', game.lower()).split() if len(w) > 1]
    if words and sum(w in t for w in words) < max(1, round(len(words) * 0.6)):
        return -99  # اسم اللعبة لازم يبان في عنوان الفيديو
    s = 4 * ('trailer' in t) + 2 * ('official' in t) + (1 if any(k in t for k in ('launch', 'story', 'reveal', 'cinematic', 'announce')) else 0)
    return s - 5 * any(b in t for b in BAD)


def search(game, key):
    q = urllib.parse.urlencode({'part': 'snippet', 'q': game + ' official trailer', 'type': 'video', 'videoEmbeddable': 'true',
                                'videoSyndicated': 'true', 'maxResults': 8, 'key': key})
    with urllib.request.urlopen('https://www.googleapis.com/youtube/v3/search?' + q, timeout=30) as r:
        return json.load(r).get('items', [])


def main(key, xlsx=R / 'Genres.xlsx', log=R / 'trailers_log.csv', search_fn=search):
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(min_row=2) if r[0].value and not (r[3].value or '').strip()]
    print(f'{len(rows)} لعبة محتاجة فيديو')
    done = failed = 0
    new = not Path(log).exists()
    with open(log, 'a', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        if new: w.writerow(['اللعبة', 'عنوان الفيديو', 'القناة', 'اللينك'])
        for r in rows:
            game = str(r[0].value).strip()
            try:
                items = search_fn(game, key)
            except urllib.error.HTTPError as e:
                body = e.read().decode('utf-8', 'ignore')
                if e.code == 403 and 'quota' in body.lower():
                    print('الكوتا اليومية خلصت — اتحفظ اللي اتعمل، كمّل بكرة.')
                else:
                    print('خطأ من يوتيوب:', e.code, body[:200])
                break
            best = max(items, key=lambda i: score(i['snippet']['title'], game), default=None)
            if best and score(best['snippet']['title'], game) >= 4:
                vid = best['id']['videoId']
                r[3].value = 'https://www.youtube.com/watch?v=' + vid
                w.writerow([game, best['snippet']['title'], best['snippet']['channelTitle'], r[3].value])
                done += 1
            else:
                w.writerow([game, '(مفيش نتيجة مناسبة)', '', ''])
                failed += 1
            if (done + failed) % 10 == 0: wb.save(xlsx)
    wb.save(xlsx)
    print(f'تم: {done} | ملقيتش تريلر مناسب: {failed} — راجع trailers_log.csv وبعدها شغّل python update.py')


if __name__ == '__main__':
    k = (sys.argv[1] if len(sys.argv) > 1 else os.environ.get('YT_KEY', '')).strip()
    sys.exit(main(k) if k else 'محتاج API key:  python fetch_trailers.py YOUR_API_KEY')
